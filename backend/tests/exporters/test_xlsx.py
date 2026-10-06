import io

from openpyxl import load_workbook

from accordance.exporters.coverage_matrix import MatrixColumn, build_coverage_matrix
from accordance.exporters.xlsx_exporter import to_xlsx_bytes
from accordance.kb.schema import Disclosure

HEADER_ROW = 4
FIRST_DATA = 5


def _disc(disc_id, standard, title):
    return Disclosure(
        id=disc_id, standard=standard, title=title, category="topic",
        requirement_text="x", required_elements=[{"id": "e", "desc": "d"}],
        retrieval_queries=["q"], suggested_fix_template="fix",
    )


def _matrix():
    kb = {
        "2-2": _disc("2-2", "GRI 2: General Disclosures 2021", "Entities"),
        "305-3": _disc("305-3", "GRI 305: Emissions 2016", "Other indirect"),
        "303-3": _disc("303-3", "GRI 303: Water 2018", "Water withdrawal"),
    }
    col = MatrixColumn(label="Acme", score_by_id={"2-2": 5, "305-3": 0, "303-3": None})
    return build_coverage_matrix(kb, [col])


def test_xlsx_has_title_block_and_header():
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(_matrix()))).active
    assert ws.title == "GRI Coverage"
    assert ws["A1"].value == "GRI Disclosure Coverage"
    assert ws["A1"].font.bold is True
    assert ws["A2"].value.startswith("Generated ")
    assert "Score key" in ws["A3"].value
    assert [ws.cell(HEADER_ROW, c).value for c in range(1, 5)] == [
        "Standard",
        "Code",
        "Indicator",
        "Acme",
    ]
    assert ws.cell(HEADER_ROW, 1).font.bold is True


def test_xlsx_score_cells():
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(_matrix()))).active
    by_code = {
        ws.cell(r, 2).value: ws.cell(r, 4).value
        for r in range(FIRST_DATA, ws.max_row + 1)
    }
    assert by_code["GRI 2-2"] == 5
    assert by_code["GRI 305-3"] == "N/A"
    assert by_code["GRI 303-3"] is None


def test_xlsx_grade_cells_color_coded():
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(_matrix()))).active
    rows = {
        ws.cell(r, 2).value: r for r in range(FIRST_DATA, ws.max_row + 1)
    }
    graded = ws.cell(rows["GRI 2-2"], 4)
    assert graded.fill.fgColor.rgb == "00DCF5E7"
    assert graded.font.bold is True
    na = ws.cell(rows["GRI 305-3"], 4)
    assert na.fill.fgColor.rgb == "00F1F1F2"
    assert na.font.italic is True


def test_xlsx_footer_has_live_formulas():
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(_matrix()))).active
    labels = {ws.cell(r, 3).value: r for r in range(FIRST_DATA, ws.max_row + 1)}
    assert "Total score" in labels
    assert "Disclosures checked" in labels
    assert "Coverage score" in labels
    tot, chk, cov = labels["Total score"], labels["Disclosures checked"], labels["Coverage score"]
    assert ws.cell(tot, 4).value == f"=SUM(D{FIRST_DATA}:D{tot - 1})"
    assert ws.cell(chk, 4).value == f"=COUNT(D{FIRST_DATA}:D{tot - 1})"
    assert ws.cell(cov, 4).value == f'=IF(D{chk}=0,"",D{tot}/(D{chk}*5))'
    assert ws.cell(cov, 4).number_format == "0.0%"
    assert ws.cell(tot, 4).font.bold is True


def test_xlsx_score_cells_integer_format_and_autofilter():
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(_matrix()))).active
    graded_row = next(
        r for r in range(FIRST_DATA, ws.max_row + 1) if ws.cell(r, 2).value == "GRI 2-2"
    )
    assert ws.cell(graded_row, 4).number_format == "0"
    assert ws.column_dimensions["C"].width == 50
    assert ws.auto_filter.ref == f"A{HEADER_ROW}:D{ws.max_row - 3}"
    assert ws.freeze_panes == "D5"
