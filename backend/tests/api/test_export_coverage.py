import io

import fitz  # PyMuPDF
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from accordance.auth.passwords import hash_password
from accordance.config import get_settings
from accordance.db import connection as db_conn
from accordance.main import create_app

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _login(client):
    with db_conn() as conn:
        conn.execute(
            "INSERT INTO users (username, password_hash) VALUES (%s, %s) "
            "ON CONFLICT (username) DO NOTHING",
            ("tester", hash_password("pw")),
        )
    client.post("/api/auth/login", json={"username": "tester", "password": "pw"})


def _seed_user(username="tester"):
    """Insert user (idempotent), return their id."""
    with db_conn() as conn:
        conn.execute(
            "INSERT INTO users (username, password_hash) VALUES (%s, %s) "
            "ON CONFLICT (username) DO NOTHING",
            (username, hash_password("pw")),
        )
        return conn.execute("SELECT id FROM users WHERE username=%s", (username,)).fetchone()["id"]


def _env(monkeypatch, tmp_path):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("LLM_MODEL", "fake:fake")
    monkeypatch.setenv("EMBEDDING_MODEL", "fake:fake")
    get_settings.cache_clear()


def _add_report(conn, report_id, name, created_at="2024-01-01 00:00:00", created_by=None):
    conn.execute(
        "INSERT INTO reports (id, name, created_at, created_by) VALUES (%s, %s, %s, %s) "
        "ON CONFLICT (id) DO NOTHING",
        (report_id, name, created_at, created_by),
    )


def _add_run(conn, run_id, report_id, version, status="completed", created_by=None):
    conn.execute(
        "INSERT INTO runs (id, report_id, version_number, kind, pdf_filename, "
        "pdf_sha256, pdf_path, status, created_by) VALUES (%s, %s, %s, 'initial', 'x.pdf', 'sha', 'x.pdf', %s, %s)",
        (run_id, report_id, version, status, created_by),
    )


def _add_finding(conn, run_id, disclosure_id, status, score=None):
    # score=None -> legacy row; effective_score back-maps the status
    # (covered->5, partial->3, missing->1). Pass score explicitly to test the
    # 0-5 path directly (e.g. score=0 for N/A).
    conn.execute(
        "INSERT INTO findings (run_id, disclosure_id, standard, status, score, note, "
        "elements_json, suggested_fix) VALUES (%s, %s, 'std', %s, %s, '', '[]', '')",
        (run_id, disclosure_id, status, score),
    )


def _add_correction(conn, run_id, disclosure_id, corrected_score, superseded=False):
    conn.execute(
        "INSERT INTO assessor_corrections "
        "(run_id, disclosure_id, standard, corrected_score, rationale, reviewer, "
        " superseded) VALUES (%s, %s, 'std', %s, 'assessor note', 'assessor', %s)",
        (run_id, disclosure_id, corrected_score, superseded),
    )


def test_export_reflects_live_assessor_correction(monkeypatch, tmp_path):
    """A saved assessor correction must override the agent score in the export,
    mirroring the UI's effectiveGrade (correction is authoritative)."""
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, created_by=uid)
        _add_finding(conn, "r1", "2-1", "partial", score=3)  # agent graded 3
        _add_correction(conn, "r1", "2-1", corrected_score=5)  # assessor corrected to 5

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/runs/r1/export/coverage")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == "GRI 2-1")
    assert ws.cell(row, 4).value == 5  # corrected score, NOT the agent's 3


def test_export_ignores_superseded_correction(monkeypatch, tmp_path):
    """Only the live (superseded=0) correction counts; an old one is ignored."""
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, created_by=uid)
        _add_finding(conn, "r1", "2-1", "partial", score=3)
        _add_correction(conn, "r1", "2-1", corrected_score=5, superseded=True)  # replaced

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/runs/r1/export/coverage")
    ws = load_workbook(io.BytesIO(resp.content)).active
    row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == "GRI 2-1")
    assert ws.cell(row, 4).value == 3  # superseded ignored -> agent's 3


def test_export_correction_to_zero_renders_na(monkeypatch, tmp_path):
    """A correction down to 0 must export as N/A (0 is a valid corrected_score)."""
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, created_by=uid)
        _add_finding(conn, "r1", "2-1", "covered", score=5)
        _add_correction(conn, "r1", "2-1", corrected_score=0)  # assessor marks N/A

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/runs/r1/export/coverage")
    ws = load_workbook(io.BytesIO(resp.content)).active
    row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == "GRI 2-1")
    assert ws.cell(row, 4).value == "N/A"


def test_per_report_export_has_one_column_per_completed_version(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, created_by=uid)
        _add_run(conn, "r2", "rep1", 2, created_by=uid)
        _add_finding(conn, "r1", "2-1", "missing")  # -> 1
        _add_finding(conn, "r2", "2-1", "covered")  # -> 5

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/reports/rep1/export/coverage")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX_CT
    assert "coverage-Acme-2024.xlsx" in resp.headers["content-disposition"]

    ws = load_workbook(io.BytesIO(resp.content)).active
    assert [c.value for c in ws[4]] == ["Standard", "Code", "Indicator", "v1", "v2"]
    row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == "GRI 2-1")
    assert ws.cell(row, 4).value == 1  # v1 missing -> 1
    assert ws.cell(row, 5).value == 5  # v2 covered -> 5


def test_export_options_lists_completed_versions_per_report(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme", "2024-01-01 00:00:00", created_by=uid)
        _add_report(conn, "rep2", "Globex", "2024-01-02 00:00:00", created_by=uid)
        _add_report(conn, "rep3", "NoRuns", "2024-01-03 00:00:00", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, created_by=uid)
        _add_run(conn, "r2", "rep1", 2, created_by=uid)
        _add_run(conn, "r3", "rep2", 1, created_by=uid)
        _add_run(
            conn, "r4", "rep3", 1, status="failed", created_by=uid
        )  # no completed run -> omitted

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/reports/export/options")
    assert resp.status_code == 200
    by_id = {d["report_id"]: d for d in resp.json()}
    # rep3 has no completed run -> not offered
    assert set(by_id) == {"rep1", "rep2"}
    # versions newest-first
    assert [v["version_number"] for v in by_id["rep1"]["versions"]] == [2, 1]
    assert by_id["rep1"]["versions"][0]["run_id"] == "r2"
    assert by_id["rep2"]["versions"] == [{"run_id": "r3", "version_number": 1}]


def test_selected_export_builds_one_column_per_run_in_order(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme", created_by=uid)
        _add_report(conn, "rep2", "Globex", created_by=uid)
        _add_run(conn, "r1", "rep1", 2, created_by=uid)
        _add_run(conn, "r2", "rep2", 1, created_by=uid)
        _add_finding(conn, "r1", "2-1", "covered")  # -> 5
        _add_finding(conn, "r2", "2-1", "partial")  # -> 3

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/reports/export/coverage?runs=r1,r2")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    # one column per selected run, in the given order, labelled "{name} v{n}"
    assert [c.value for c in ws[4]] == ["Standard", "Code", "Indicator", "Acme v2", "Globex v1"]
    row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == "GRI 2-1")
    assert ws.cell(row, 4).value == 5
    assert ws.cell(row, 5).value == 3


def test_selected_export_single_run_uses_report_filename(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 2, created_by=uid)
        _add_finding(conn, "r1", "2-1", "covered")

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/reports/export/coverage?runs=r1")
    assert resp.status_code == 200
    assert "coverage-Acme-2024-v2.xlsx" in resp.headers["content-disposition"]


def test_selected_export_requires_at_least_one_run(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    client = TestClient(create_app())
    _login(client)
    assert client.get("/api/reports/export/coverage").status_code == 400
    assert client.get("/api/reports/export/coverage?runs=").status_code == 400


def test_selected_export_rejects_noncompleted_or_unknown_run(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, status="failed", created_by=uid)

    client = TestClient(create_app())
    _login(client)
    assert client.get("/api/reports/export/coverage?runs=r1").status_code == 400
    assert client.get("/api/reports/export/coverage?runs=ghost").status_code == 404


def test_export_renders_na_for_zero_score_and_footer_formulas(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, created_by=uid)
        _add_finding(conn, "r1", "2-1", "missing", score=0)  # explicit N/A
        _add_finding(conn, "r1", "2-2", "covered", score=5)

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/runs/r1/export/coverage")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    cells = {ws.cell(r, 2).value: ws.cell(r, 4).value for r in range(5, ws.max_row + 1)}
    assert cells["GRI 2-1"] == "N/A"  # score 0 -> N/A
    assert cells["GRI 2-2"] == 5
    # Footer formulas present (Total / Checked exclude the N/A row).
    footer = {ws.cell(r, 3).value: ws.cell(r, 4).value for r in range(5, ws.max_row + 1)}
    assert footer["Total score"].startswith("=SUM(")
    assert footer["Disclosures checked"].startswith("=COUNT(")
    assert footer["Coverage score"].startswith("=IF(")


def test_per_report_404_when_no_completed_runs(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme", created_by=uid)
        _add_run(conn, "rep1", "rep1", 1, status="failed", created_by=uid)

    client = TestClient(create_app())
    _login(client)
    assert client.get("/api/reports/rep1/export/coverage").status_code == 404


def test_single_run_export_has_one_column_labeled_report_and_version(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 2, created_by=uid)
        _add_finding(conn, "r1", "2-1", "partial")  # -> 3

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/runs/r1/export/coverage")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX_CT
    assert "coverage-Acme-2024-v2.xlsx" in resp.headers["content-disposition"]

    ws = load_workbook(io.BytesIO(resp.content)).active
    assert [c.value for c in ws[4]] == ["Standard", "Code", "Indicator", "Acme 2024 v2"]
    # Only the tested disclosure is present (footer rows have an empty Code column).
    codes = {ws.cell(r, 2).value for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value}
    assert codes == {"GRI 2-1"}
    row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == "GRI 2-1")
    assert ws.cell(row, 4).value == 3  # partial -> 3


def test_single_run_export_404_when_run_missing(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    client = TestClient(create_app())
    _login(client)
    assert client.get("/api/runs/nope/export/coverage").status_code == 404


def test_analysis_export_returns_pdf_for_completed_run(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme 2024", created_by=uid)
        _add_run(conn, "r1", "rep1", 2, created_by=uid)
        _add_finding(conn, "r1", "2-1", "partial", score=3)
        conn.execute(
            "INSERT INTO chunks (run_id, page, text) VALUES ('r1', 12, 'x')"
        )

    client = TestClient(create_app())
    _login(client)
    resp = client.get("/api/runs/r1/export/analysis")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert "analysis-Acme-2024-v2.pdf" in resp.headers["content-disposition"]
    assert resp.content.startswith(b"%PDF")
    text = fitz.open(stream=resp.content, filetype="pdf")[0].get_text()
    assert "Acme 2024" in text
    assert "r1" in text
    assert "GRI 2: General Disclosures 2021" in text


def test_analysis_export_rejects_incomplete_run(monkeypatch, tmp_path):
    _env(monkeypatch, tmp_path)
    uid = _seed_user()
    with db_conn() as conn:
        _add_report(conn, "rep1", "Acme", created_by=uid)
        _add_run(conn, "r1", "rep1", 1, status="failed", created_by=uid)
        _add_run(conn, "r2", "rep1", 2, status="judging", created_by=uid)

    client = TestClient(create_app())
    _login(client)
    assert client.get("/api/runs/r1/export/analysis").status_code == 400
    assert client.get("/api/runs/r2/export/analysis").status_code == 400
    assert client.get("/api/runs/nope/export/analysis").status_code == 404
