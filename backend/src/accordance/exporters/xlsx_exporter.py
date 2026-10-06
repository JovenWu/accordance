import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from accordance.exporters.coverage_matrix import CoverageMatrix

_LABEL_COLS = 3

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, color="18181B", size=16)
_META_FONT = Font(color="71717A", size=10)
_LEGEND_FONT = Font(color="71717A", size=9, italic=True)
_LABEL_FONT = Font(color="18181B", size=10)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_SCORE_FMT = "0"
_COVERAGE_FMT = "0.0%"
_NA = "N/A"

_GRADE_STYLE: dict[int, tuple[str, str]] = {
    5: ("DCF5E7", "0F7A43"),
    4: ("E9F3D9", "527A1D"),
    3: ("FCEFD3", "8A5A00"),
    2: ("FDE7D6", "B04A0C"),
    1: ("FFDCDD", "B4232A"),
}
_NA_STYLE = ("F1F1F2", "6B7280")
_ZEBRA_FILL = PatternFill("solid", fgColor="F7F7F8")
_FOOTER_FILL = PatternFill("solid", fgColor="F4F4F5")

_THIN = Side(style="thin", color="D4D4D8")
_CELL_BORDER = Border(bottom=_THIN)
_FOOTER_TOP = Border(top=Side(style="thin", color="9CA3AF"), bottom=_THIN)

_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")

_TITLE_ROW = 1
_META_ROW = 2
_LEGEND_ROW = 3
_HEADER_ROW = 4
_FIRST_DATA = 5


def _safe(value):
    """Neutralize spreadsheet formula injection in a user-controlled string.

    Report names (column labels) and KB-derived labels (standard / code /
    indicator) flow into cells, and these exports are shared with assessors, so
    a value beginning with a formula trigger is prefixed with a single quote to
    force literal text. Non-str (and empty) values pass through unchanged — the
    server-authored footer formulas are written separately and stay live.
    """
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGERS):
        return "'" + value
    return value


def to_xlsx_bytes(matrix: CoverageMatrix) -> bytes:
    """Render a CoverageMatrix as a polished .xlsx worksheet (bytes).

    Layout: title and a generated-on line, a score legend, then the matrix —
    header row, one row per disclosure, and a live-formula footer. Each data
    cell holds the disclosure's 0-5 grade color-coded like the app's score
    pills (0 rendered as "N/A"); blanks are disclosures a column did not judge.
    A three-row footer per column carries LIVE Excel formulas — not
    precomputed values:

      Total score          =SUM(<grades>)        (1-5 cells; N/A text + blanks ignored)
      Disclosures checked  =COUNT(<grades>)       (applicable 1-5 cells only)
      Coverage score       =Total/(Checked*5)     (shown as %; matches the app's Completeness %)

    N/A and un-judged disclosures never drag the coverage score down — the
    denominator counts only applicable (1-5) disclosures, mirroring the
    dashboard's completeness metric.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "GRI Coverage"

    n_label = _LABEL_COLS
    n_cols = n_label + len(matrix.column_labels)
    last_letter = get_column_letter(n_cols)

    ws.merge_cells(f"A{_TITLE_ROW}:{last_letter}{_TITLE_ROW}")
    t = ws.cell(_TITLE_ROW, 1, "GRI Disclosure Coverage")
    t.font = _TITLE_FONT
    t.alignment = _LEFT
    ws.row_dimensions[_TITLE_ROW].height = 28

    ws.merge_cells(f"A{_META_ROW}:{last_letter}{_META_ROW}")
    m = ws.cell(
        _META_ROW,
        1,
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} · "
        f"{len(matrix.column_labels)} "
        f"report column{'s' if len(matrix.column_labels) != 1 else ''}",
    )
    m.font = _META_FONT
    m.alignment = _LEFT
    ws.row_dimensions[_META_ROW].height = 16

    ws.merge_cells(f"A{_LEGEND_ROW}:{last_letter}{_LEGEND_ROW}")
    lg = ws.cell(
        _LEGEND_ROW,
        1,
        "Score key — 5 Complete · 4 Substantial · 3 Partial · 2 Minimal · "
        "1 Not found · N/A Not applicable · blank Not assessed",
    )
    lg.font = _LEGEND_FONT
    lg.alignment = _LEFT
    ws.row_dimensions[_LEGEND_ROW].height = 16

    for col, label in enumerate(
        ["Standard", "Code", "Indicator", *matrix.column_labels], start=1
    ):
        cell = ws.cell(_HEADER_ROW, col, _safe(label))
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER_WRAP if col > n_label else _LEFT_WRAP
    ws.row_dimensions[_HEADER_ROW].height = 32

    for i, row in enumerate(matrix.rows):
        r = _FIRST_DATA + i
        for col, value in enumerate(
            (row.standard, row.code, row.indicator), start=1
        ):
            cell = ws.cell(r, col, _safe(value))
            cell.font = _LABEL_FONT
            cell.alignment = _CENTER if col == 2 else _LEFT_WRAP
            cell.border = _CELL_BORDER
            if i % 2 == 1:
                cell.fill = _ZEBRA_FILL
        for j, v in enumerate(row.values):
            c = ws.cell(r, n_label + 1 + j)
            c.border = _CELL_BORDER
            if v is None:
                continue
            if v == 0:
                c.value = _NA
                fill, fg = _NA_STYLE
                c.font = Font(color=fg, size=10, italic=True)
            else:
                c.value = v
                c.number_format = _SCORE_FMT
                fill, fg = _GRADE_STYLE[v]
                c.font = Font(color=fg, size=10, bold=True)
            c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = _CENTER
        ws.row_dimensions[r].height = 18
    last_data = _FIRST_DATA + len(matrix.rows) - 1
    have_data = len(matrix.rows) > 0

    total_row = last_data + 1
    checked_row = last_data + 2
    coverage_row = last_data + 3
    for rr, label in (
        (total_row, "Total score"),
        (checked_row, "Disclosures checked"),
        (coverage_row, "Coverage score"),
    ):
        cell = ws.cell(rr, n_label, label)
        cell.font = Font(bold=True, size=10)
        cell.alignment = _LEFT
        cell.fill = _FOOTER_FILL
        cell.border = _FOOTER_TOP if rr == total_row else _CELL_BORDER

    for j in range(len(matrix.column_labels)):
        col = n_label + 1 + j
        letter = get_column_letter(col)
        if have_data:
            rng = f"{letter}{_FIRST_DATA}:{letter}{last_data}"
            ws.cell(total_row, col).value = f"=SUM({rng})"
            ws.cell(checked_row, col).value = f"=COUNT({rng})"
            ws.cell(coverage_row, col).value = (
                f'=IF({letter}{checked_row}=0,"",'
                f"{letter}{total_row}/({letter}{checked_row}*5))"
            )
        else:
            ws.cell(total_row, col).value = 0
            ws.cell(checked_row, col).value = 0
            ws.cell(coverage_row, col).value = ""
        ws.cell(total_row, col).number_format = _SCORE_FMT
        ws.cell(checked_row, col).number_format = _SCORE_FMT
        ws.cell(coverage_row, col).number_format = _COVERAGE_FMT
        for rr in (total_row, checked_row, coverage_row):
            cell = ws.cell(rr, col)
            cell.alignment = _CENTER
            cell.font = Font(bold=True, size=10)
            cell.fill = _FOOTER_FILL
            cell.border = (
                _FOOTER_TOP if rr == total_row else _CELL_BORDER
            )
    for col in range(1, n_label):
        for rr in (total_row, checked_row, coverage_row):
            ws.cell(rr, col).fill = _FOOTER_FILL
        ws.cell(total_row, col).border = _FOOTER_TOP

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 50
    for col in range(n_label + 1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    filter_last = last_data if have_data else _HEADER_ROW
    ws.auto_filter.ref = f"A{_HEADER_ROW}:{last_letter}{filter_last}"

    ws.freeze_panes = f"D{_FIRST_DATA}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
